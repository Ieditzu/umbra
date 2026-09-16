"""Build the editable RoSEF dossier and the 80 x 120 cm working poster.

Run in a document environment with python-docx and reportlab installed.
The dossier uses standard_business_brief with memo_masthead; title override
is 26 pt, subtitle 12 pt. Poster typography uses embedded Noto Sans.
"""

import json
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "competition"
RESULTS = json.loads((OUT / "software-results.json").read_text())
LABELS = {"directional_target": "Lumină direcțională, țintă prezentă",
          "directional_empty": "Lumină direcțională, fără țintă",
          "flat_target": "Lumină difuză, țintă prezentă"}


def results_text():
    text = ["Trei demonstrații sintetice: fiecare are 120 de cadre la 10 cadre/s, seed 7. "
            "Rezultate generate la " + RESULTS["generated_at"][:10] + "."]
    for case in RESULTS["cases"]:
        metrics = case["metrics"]
        error = metrics["position_error_norm"]
        false_alert = metrics["false_alert_rate"]
        text.append(f"{LABELS[case['case']]}: {case['state']}; scor de încredere {case['confidence']:.4f}; "
                    f"F1 {metrics['f1']:.4f}; TP={metrics['true_positive']}, FP={metrics['false_positive']}, "
                    f"FN={metrics['false_negative']}, TN={metrics['true_negative']}. "
                    f"Eroare normalizată: {error if error is not None else 'indisponibilă'}. "
                    f"Rata alarmelor false: {false_alert if false_alert is not None else 'indisponibilă' }.")
    text.append("Reproducere: scripts/validate_synthetic.py. Date și versiuni: software-results.json. "
                "Eroarea de poziție este raportată pentru indicatorul brut, inclusiv când rezultatul global este refuzat.")
    return "\n\n".join(text)


def dossier():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    for name, size, color, before, after in (
        ("Normal", 11, "222222", 0, 6), ("Title", 26, "222222", 0, 8),
        ("Subtitle", 12, "555555", 0, 6), ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6), ("Heading 3", 12, "1F4D78", 8, 4),
        ("Header", 9, "666666", 0, 0), ("Footer", 9, "666666", 0, 0),
    ):
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10
        style.paragraph_format.widow_control = True
        style.paragraph_format.keep_with_next = name.startswith("Heading")
    section.header.paragraphs[0].text = "UMBRA  /  Dosar de lucru RoSEF  /  Ediție de confirmat"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Document pregătitor · ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    content = (OUT / "dossier.ro.md").read_text().replace("{{SOFTWARE_RESULTS}}", results_text())
    for block in content.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if block == "<!-- PAGE -->":
            doc.add_page_break()
        elif block.startswith("### "):
            doc.add_paragraph(block[4:], "Heading 2")
        elif block.startswith("## "):
            doc.add_paragraph(block[3:], "Heading 1")
        elif block.startswith("# "):
            doc.add_paragraph(block[2:], "Title")
        else:
            doc.add_paragraph(block)
    doc.core_properties.title = "Umbra — dosar de lucru RoSEF"
    doc.core_properties.author = ""
    doc.core_properties.subject = "Document pregătitor; datele participanților și validarea în teren lipsesc"
    doc.save(OUT / "umbra-dosar.docx")
    # Structural check supplements, but does not replace, Word rendering.
    reopened = Document(OUT / "umbra-dosar.docx")
    assert len(reopened.paragraphs) == len(doc.paragraphs)
    assert "{{SOFTWARE_RESULTS}}" not in "\n".join(p.text for p in reopened.paragraphs)


def poster():
    pdfmetrics.registerFont(TTFont("Noto", "/usr/share/fonts/noto/NotoSans-Regular.ttf"))
    pdfmetrics.registerFont(TTFont("NotoBold", "/usr/share/fonts/noto/NotoSans-Bold.ttf"))
    width, height = 80 * cm, 120 * cm
    pdf = canvas.Canvas(str(OUT / "umbra-poster.pdf"), pagesize=(width, height))
    pdf.setTitle("Umbra — poster de lucru RoSEF, dovezi sintetice")
    pdf.setFillColor(HexColor("#FAF9F5"))
    pdf.rect(0, 0, width, height, fill=1, stroke=0)
    ink, accent, muted = "#162A33", "#226759", "#536168"
    margin = 100

    def para(text, x, top, available, size=34, color=ink, bold=False):
        style = ParagraphStyle("poster", fontName="NotoBold" if bold else "Noto", fontSize=size,
                               leading=size * 1.32, textColor=HexColor(color))
        p = Paragraph(escape(text), style)
        _, used = p.wrap(available, height)
        assert top + used < height - 50, (text, top, used)
        p.drawOn(pdf, x, height - top - used)
        return top + used

    para("UMBRA", margin, 95, width - 2 * margin, 116, bold=True)
    para("Mișcare ascunsă. Indicii în lumină.", margin, 270, width - 2 * margin, 57, bold=True)
    para("Kit local pentru studiul variațiilor de lumină lângă un colț", margin, 375, width - 2 * margin, 38)
    para("Autor(i), instituție, coordonator: DE COMPLETAT  ·  Ediție RoSEF: DE CONFIRMAT", margin, 465,
         width - 2 * margin, 25, muted)
    pdf.setFillColor(HexColor(accent))
    pdf.rect(margin, height - 635, width - 2 * margin, 88, fill=1, stroke=0)
    para("PROTOTIP SOFTWARE · REZULTATE SINTETICE · VALIDARE ÎN TEREN NECESARĂ", margin + 25, 570,
         width - 2 * margin - 50, 27, "#FFFFFF", True)

    col = (width - 2 * margin - 90) / 2
    right = margin + col + 90
    top = 720
    para("ÎNTREBAREA", margin, top, col, 34, accent, True)
    para("Poate o cameră fixă să extragă indicii despre o țintă ascunsă din lumina vizibilă lângă un colț?",
         margin, top + 75, col, 39)
    para("METODA", right, top, col, 34, accent, True)
    para("Calibrare fără țintă; regiune automată; normalizarea expunerii; variații temporale; estimare și refuz.",
         right, top + 75, col, 37)

    top = 1110
    para("CE OBSERVĂ CAMERA", margin, top, width - 2 * margin, 34, accent, True)
    # Schematic: opaque corner, visible patch, main camera, hidden target.
    origin_y = height - 1650
    pdf.setFillColor(HexColor("#DCDDD6"))
    pdf.rect(margin + 50, origin_y, 510, 360, fill=1, stroke=0)
    pdf.setStrokeColor(HexColor(ink))
    pdf.setLineWidth(9)
    pdf.lines([(margin + 560, origin_y, margin + 560, origin_y + 360),
               (margin + 50, origin_y, margin + 560, origin_y)])
    pdf.setFillColor(HexColor("#B8D9C8"))
    pdf.ellipse(margin + 590, origin_y + 20, margin + 860, origin_y + 155, fill=1, stroke=0)
    pdf.setFillColor(HexColor(accent))
    pdf.circle(margin + 705, origin_y + 320, 34, fill=1, stroke=0)
    pdf.setFillColor(HexColor(ink))
    pdf.rect(margin + 190, origin_y - 150, 85, 55, fill=1, stroke=0)
    pdf.setStrokeColor(HexColor(muted))
    pdf.setLineWidth(3)
    pdf.line(margin + 275, origin_y - 120, margin + 725, origin_y + 80)
    para("Țintă ascunsă", margin + 755, 1280, 400, 25)
    para("Perete", margin + 160, 1390, 350, 28)
    para("Zonă analizată", margin + 610, 1660, 430, 25)
    para("Cameră principală", margin + 50, 1820, 450, 27)
    para("Schema este conceptuală. Camera înregistrează suprafața vizibilă; o a doua cameră furnizează numai etichete de referință.",
         right + 100, 1285, col - 100, 36)
    para("Poziția estimată este un indicator derivat din imagine, încă necalibrat fizic.",
         right + 100, 1620, col - 100, 34, muted)

    top = 1980
    para("TREI VERIFICĂRI SINTETICE", margin, top, width - 2 * margin, 34, accent, True)
    para("Seed 7 · 120 cadre/caz · 10 cadre/s · aceleași date generate determinist", margin, top + 66,
         width - 2 * margin, 28, muted)
    states = {"strong_motion": "Mișcare puternică", "no_reliable_signal": "Fără semnal fiabil"}
    for index, case in enumerate(RESULTS["cases"]):
        y = top + 150 + index * 115
        pdf.setStrokeColor(HexColor("#C8CEC8"))
        pdf.setLineWidth(1)
        pdf.line(margin, height - y + 15, width - margin, height - y + 15)
        para(LABELS[case["case"]], margin, y, 950, 30)
        para(states[case["state"]], margin + 990, y, 600, 30, bold=True)
        para(f"F1 {case['metrics']['f1']:.3f}", width - margin - 270, y, 270, 30)
    para("F1 este calculat pe cadre, înaintea refuzului global. În cazul fără țintă: 3/120 cadre peste prag (2,5%). "
         "Cadrele sunt corelate; acestea nu sunt încercări independente și nu demonstrează performanță în exterior.",
         margin, 2510, width - 2 * margin, 30)

    para("CONTRIBUȚIA PROPUSĂ", margin, 2710, col, 32, accent, True)
    para("Un flux local, reproductibil, cu calibrare, explicații, rezultate salvate și refuz explicit. "
         "Aportul individual și asistența AI sunt de declarat în dosar.", margin, 2780, col, 33)
    para("CE URMEAZĂ", right, 2710, col, 32, accent, True)
    para("Încercări controlate la trei colțuri, țintă prezentă/absentă, etichete revizuite și evaluare pe colțuri "
         "nefolosite la reglaj. Ipoteza rămâne de testat.", right, 2780, col, 33)
    para("Referință: Bouman și colab., Turning Corners Into Cameras: Principles and Methods, ICCV 2017. "
         "Principiul fizic este anterior proiectului. Dovezi software: software-results.json; reproducere: scripts/validate_synthetic.py.",
         margin, 3140, width - 2 * margin, 23, muted)
    pdf.showPage()
    pdf.save()


if __name__ == "__main__":
    dossier()
    poster()
    print(f"Created {OUT / 'umbra-dosar.docx'} and {OUT / 'umbra-poster.pdf'}")
