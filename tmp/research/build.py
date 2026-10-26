from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/maneacezar/Documents/GitHub/umbra')
source = (ROOT / 'tmp/research/paper.md').read_text()
out = ROOT / 'output/research'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.9)
sec.left_margin, sec.right_margin = Cm(2.1), Cm(2.1)
sec.header_distance, sec.footer_distance = Cm(.8), Cm(.8)
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(11.5)
normal.paragraph_format.line_spacing = 1.09
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True
for name, size in [('Title',23),('Subtitle',14),('Heading 1',14),('Heading 2',12)]:
    st = doc.styles[name]
    st.font.name = 'Times New Roman'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0,0,0)
    st.font.bold = name != 'Subtitle'
    st.paragraph_format.space_before = Pt(12 if name.startswith('Heading') else 0)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.keep_with_next = True
header = sec.header.paragraphs[0]
header.text = 'UMBRA SmartRoute    |    Studiu științific    |    RoSEF 2026'
header.style = doc.styles['Normal']
header.runs[0].font.size = Pt(9)
header.paragraph_format.space_after = Pt(0)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run('UMBRA SmartRoute  •  ')
r.font.size = Pt(9)
field = OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def table(rows):
    cells = [r.strip().strip('|').split('|') for r in rows]
    t = doc.add_table(rows=1, cols=len(cells[0]))
    t.style = 'Table Grid'
    for i, vals in enumerate(cells):
        row = t.rows[0] if i == 0 else t.add_row()
        for cell, val in zip(row.cells, vals):
            cell.text = val.strip()
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                for rr in p.runs:
                    rr.font.size = Pt(9.5)
                    rr.bold = i == 0
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trpr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

lines = source.splitlines()
i=0
while i < len(lines):
    line=lines[i].strip()
    if not line:
        i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            rows.append(lines[i]);i+=1
        table(rows);continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:],'Title')
    elif line.startswith('## '):
        p=doc.add_paragraph(line[3:],'Subtitle' if line.startswith('## Optimizarea') else 'Heading 1')
    elif line.startswith('### '):
        p=doc.add_paragraph(line[4:],'Heading 2')
    elif line.startswith('EQ '):
        p=doc.add_paragraph(line[3:]);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(9)
        for rr in p.runs:
            rr.font.name='Cambria';rr.font.size=Pt(11)
    else:
        p=doc.add_paragraph(line)
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        if line.startswith('['):
            p.paragraph_format.line_spacing=1.03
            p.paragraph_format.space_after=Pt(7)
            for rr in p.runs:rr.font.size=Pt(10)
        if 'http' in line:
            # Allow long source URLs to wrap without changing their visible content.
            for rr in p.runs:
                rr.text = re.sub(r'https?://\S+', lambda m: m.group().replace('/','/\u200b').replace('-','-\u200b'), rr.text)
    i+=1
doc.core_properties.title = 'UMBRA SmartRoute Optimizarea traseelor pietonale umbrite'
doc.core_properties.subject = 'Studiu de fundamentare științifică și protocol experimental'
doc.core_properties.author = ''
doc.core_properties.keywords = 'UMBRA, SmartRoute, umbrire, incertitudine, RoSEF 2026'
for el in doc.styles.element.iter():
    if el.tag == qn('w:pBdr'):
        el.getparent().remove(el)
for p in doc.paragraphs:
    if p.style.name == 'Title':
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.color.rgb = RGBColor(0,0,0)
            rf = r._element.get_or_add_rPr().get_or_add_rFonts()
            for key in list(rf.attrib):
                if 'theme' in key.lower(): del rf.attrib[key]
doc.save(out/'UMBRA_SmartRoute_Studiu_Stiintific_RoSEF_2026.docx')
print('Words',len(source.split()))
print('Saved',out/'UMBRA_SmartRoute_Studiu_Stiintific_RoSEF_2026.docx')
